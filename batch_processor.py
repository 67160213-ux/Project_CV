import os
import glob
import pandas as pd
from tqdm import tqdm
from scan_and_predict import scan_sheet

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TRAIN_DIR = os.path.join(BASE_DIR, "input_images")
TEST_DIR = os.path.join(BASE_DIR, "test_image")
INPUT_DIR = TEST_DIR  # ดึงภาพทดสอบจาก test_image เป็นค่าเริ่มต้น
OUTPUT_EXCEL = os.path.join(BASE_DIR, "outputs", "results", "all_results.xlsx")

def export_styled_excel(df, output_excel):
    """Exports results with professional styling: Dashboard sheet + Styled Detailed sheet"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import datetime

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    # ---------------- 1. Dashboard Sheet ----------------
    ws_dash = wb.create_sheet(title="Dashboard_ภาพรวม")
    ws_dash.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_dash.merge_cells("A1:G2")
    t_cell = ws_dash["A1"]
    t_cell.value = "📊 รายงานสรุปผลการตรวจกระดาษคำตอบอัตโนมัติ (Automated Answer Sheet Report)"
    t_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_dash["A3"].value = f"สร้างรายงานเมื่อ: {now_str} | ตรวจทั้งหมด: {df['Image_Name'].nunique()} แผ่น | รวม {len(df)} ข้อ"
    ws_dash["A3"].font = Font(name="Segoe UI", size=10, italic=True, color="595959")

    # Ground Truth Template: ชุดที่ 1 (ประโยคภาษาอังกฤษ & Tech 20 แถว x 20 ช่อง)
    GROUND_TRUTH_TECH = {
        1:  "THE-QUICK-BROWN-FOX-",
        2:  "JUMPS-OVER-LAZY-DOG-",
        3:  "PACK-MY-BOX-WITH-TEA",
        4:  "FIVE-DOZEN-LIQUOR-JU",
        5:  "SPHINX-OF-BLACK-GEM-",
        6:  "JUDGE-MY-VOW-QUICKLY",
        7:  "BRIGHT-ZEPHYRS-BLOW-",
        8:  "COMPUTER-VISION-TEST",
        9:  "MACHINE-LEARNING-AI-",
        10: "DEEP-NEURAL-NETWORKS",
        11: "IMAGE-DATA-DETECT-OK",
        12: "PYTHON-OPENCV-ROBOT-",
        13: "SMART-OCR-RECOGNIZE-",
        14: "ACCURATE-PREDICTION-",
        15: "ARTIFICIAL-SYSTEM-OK",
        16: "HELLO-WORLD-WELCOME-",
        17: "KNOWLEDGE-IS-POWER--",
        18: "FIND-EXTRA-FINE-GOLD",
        19: "ZERO-ONE-TWO-SIX-TEN",
        20: "ABCDEFGHIJKLMNOPQRST",
    }

    GROUND_TRUTH_ALPHABET = {
        q: ("ABCDEFGHIJKLM-------" if q % 2 != 0 else "NOPQRSTUVWXYZ-------")
        for q in range(1, 21)
    }

    def get_sheet_ground_truth(sheet_records):
        tech_score = 0
        alpha_score = 0
        for r in sheet_records:
            q = int(r["Question"])
            ans = str(r["Recognized_Answers"])
            exp_t = GROUND_TRUTH_TECH.get(q, "-" * 20)
            exp_a = GROUND_TRUTH_ALPHABET.get(q, "-" * 20)
            for ca, ct, caa in zip(ans, exp_t, exp_a):
                if ca == ct:
                    tech_score += 1
                if ca == caa:
                    alpha_score += 1
        if alpha_score > tech_score:
            return GROUND_TRUTH_ALPHABET, "Alphabet (A-Z)"
        return GROUND_TRUTH_TECH, "English & Tech (ชุดที่ 1)"

    # Overall Metrics Calculation
    total_chars = len(df) * 20
    correct_chars = 0
    sheet_gt_map = {}
    for img_name, grp in df.groupby("Image_Name", sort=False):
        gt_dict, tmpl_name = get_sheet_ground_truth(grp.to_dict("records"))
        sheet_gt_map[img_name] = (gt_dict, tmpl_name)
        for _, r in grp.iterrows():
            q = int(r["Question"])
            ans = str(r["Recognized_Answers"])
            exp = gt_dict.get(q, "-" * 20)
            for ca, ce in zip(ans, exp):
                if ca == ce:
                    correct_chars += 1
    acc_total = (correct_chars / total_chars * 100) if total_chars > 0 else 0.0

    kpis = [
        ("จำนวนแผ่นที่ตรวจ", f"{df['Image_Name'].nunique()} แผ่น", "D9E1F2", "203764"),
        ("จำนวนข้อทั้งหมด", f"{len(df)} ข้อ", "D9E1F2", "203764"),
        ("จำนวนตัวอักษร", f"{total_chars:,} ตัว", "D9E1F2", "203764"),
        ("ความแม่นยำรวม", f"{acc_total:.2f}%", "E2EFDA", "375623"),
        ("ข้อที่ต้องตรวจทาน", f"{(df['Needs_Review'] == 'YES').sum()} ข้อ", "FCE4D6", "C65911"),
    ]

    cols = ["A", "B", "C", "D", "E"]
    for idx, (label, val, bg, fg) in enumerate(kpis):
        c_letter = cols[idx]
        lbl_cell = ws_dash[f"{c_letter}5"]
        lbl_cell.value = label
        lbl_cell.font = Font(name="Segoe UI", size=10, bold=True, color="595959")
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        lbl_cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

        val_cell = ws_dash[f"{c_letter}6"]
        val_cell.value = val
        val_cell.font = Font(name="Segoe UI", size=20, bold=True, color=fg)
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

    # Table per-sheet Breakdown
    ws_dash["A9"].value = "📋 สรุปเปรียบเทียบรายแผ่น (Per-Sheet Breakdown)"
    ws_dash["A9"].font = Font(name="Segoe UI", size=11, bold=True, color="1F4E78")

    dash_headers = ["ชื่อไฟล์ (Image)", "จำนวนข้อ", "จำนวนตัวอักษร", "ความแม่นยำ (%)", "Avg Confidence", "ข้อที่เตือน (Review)", "ชุดเฉลยที่ตรวจ"]
    for c_i, h in enumerate(dash_headers, 1):
        cell = ws_dash.cell(row=10, column=c_i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    row_c = 11
    for img_name, grp in df.groupby("Image_Name", sort=False):
        gt_dict, tmpl_name = sheet_gt_map.get(img_name, (GROUND_TRUTH_TECH, "English & Tech (ชุดที่ 1)"))
        t_chars = len(grp) * 20
        c_chars = 0
        for _, r_data in grp.iterrows():
            q = int(r_data["Question"])
            ans = str(r_data["Recognized_Answers"])
            exp = gt_dict.get(q, "-" * 20)
            for ca, ce in zip(ans, exp):
                if ca == ce:
                    c_chars += 1
        acc = (c_chars / t_chars * 100) if t_chars > 0 else 0.0
        avg_c = grp["Average_Confidence"].mean()
        rev_c = (grp["Needs_Review"] == "YES").sum()
        stat = "สมบูรณ์ 100%" if acc == 100.0 else tmpl_name

        row_vals = [img_name, len(grp), t_chars, f"{acc:.2f}%", f"{avg_c:.3f}", f"{rev_c} ข้อ", stat]
        for c_i, val in enumerate(row_vals, 1):
            cell = ws_dash.cell(row=row_c, column=c_i, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if c_i == 4 and acc == 100.0:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="375623")
                cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            elif c_i == 4:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="203764")
            if c_i == 7 and acc == 100.0:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="375623")
            elif c_i == 7:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="C65911")
        row_c += 1

    # ---------------- 2. Detailed Results Sheet ----------------
    ws_detail = wb.create_sheet(title="Detailed_Results")
    ws_detail.views.sheetView[0].showGridLines = True

    detail_headers = ["Image_Name", "Question", "Recognized_Answers", "Average_Confidence", "Min_Confidence", "Needs_Review"]
    for c_i, h in enumerate(detail_headers, 1):
        cell = ws_detail.cell(row=1, column=c_i, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_detail.row_dimensions[1].height = 25

    for r_idx, row_data in df.iterrows():
        r_num = r_idx + 2
        ws_detail.row_dimensions[r_num].height = 20
        bg_c = "FFFFFF" if (r_idx % 2 == 0) else "F9FAFC"

        for c_i, col_name in enumerate(detail_headers, 1):
            val = row_data[col_name]
            cell = ws_detail.cell(row=r_num, column=c_i, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            cell.fill = PatternFill(start_color=bg_c, end_color=bg_c, fill_type="solid")

            if col_name in ["Image_Name", "Question"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name == "Recognized_Answers":
                cell.font = Font(name="Consolas", size=11, bold=True, color="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name in ["Average_Confidence", "Min_Confidence"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "0.000"
            elif col_name == "Needs_Review":
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if val == "YES":
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="C00000")
                    cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                else:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="375623")
                    cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # Column Auto-width
    for ws in [ws_dash, ws_detail]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.coordinate in ["A1", "B1", "C1", "D1", "E1", "F1", "G1", "A2", "B2", "C2", "D2", "E2", "F2", "G2"]:
                    continue
                v_str = str(cell.value or "")
                max_len = max(max_len, len(v_str))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    ws_detail.freeze_panes = "A2"
    ws_detail.auto_filter.ref = f"A1:F{len(df)+1}"

    os.makedirs(os.path.dirname(output_excel), exist_ok=True)
    wb.save(output_excel)

def run_batch(input_folder=INPUT_DIR, output_excel=OUTPUT_EXCEL):
    extensions = ("*.jpg", "*.jpeg", "*.png")
    image_paths = []
    for ext in extensions:
        image_paths.extend(glob.glob(os.path.join(input_folder, ext)))
        
    print(f"พบรูปภาพใน {input_folder} ทั้งหมด: {len(image_paths)} รูป")
    if not image_paths:
        print(f"ยังไม่มีภาพในโฟลเดอร์ {input_folder} ให้ประมวลผล")
        return
        
    all_records = []
    failed_images = []
    
    for img_path in tqdm(image_paths, desc="กำลังประมวลผล"):
        filename = os.path.basename(img_path)
        try:
            # ประมวลผลแต่ละแผ่น (เซฟรูปผลลัพธ์ลง outputs/visuals/ อัตโนมัติ)
            sheet_data = scan_sheet(img_path, save_annotated=True)
            
            for q_num, chars_info in sheet_data.items():
                ans_str = "".join([c["character"] for c in chars_info])
                conf_list = [c["confidence"] for c in chars_info]
                avg_conf = sum(conf_list) / max(len(conf_list), 1)
                min_conf = min(conf_list) if conf_list else 0.0
                
                all_records.append({
                    "Image_Name": filename,
                    "Question": q_num,
                    "Recognized_Answers": ans_str,
                    "Average_Confidence": round(avg_conf, 3),
                    "Min_Confidence": round(min_conf, 3),
                    "Needs_Review": "YES" if min_conf < 0.80 else "NO"
                })
        except Exception as e:
            print(f"\n[Error] {filename}: {e}")
            failed_images.append({"Image": filename, "Error": str(e)})
            
    df = pd.DataFrame(all_records)
    export_styled_excel(df, output_excel)
    print(f"\nประมวลผลเสร็จสิ้น! บันทึกไฟล์ Excel สรุปผลลัพธ์ที่: {output_excel}")
    
    if failed_images:
        print(f"มีภาพที่มีปัญหา {len(failed_images)} รูป:")
        for fail in failed_images:
            print(f" - {fail['Image']}: {fail['Error']}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else INPUT_DIR
    run_batch(input_folder=target)
